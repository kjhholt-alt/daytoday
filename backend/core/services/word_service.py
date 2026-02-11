import logging
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

logger = logging.getLogger(__name__)

# File extensions this service can scan and extract text from
_SUPPORTED_EXTENSIONS = {'*.docx', '*.pdf'}


class WordService:
    def __init__(self, directories: list):
        self._directories = [Path(d) for d in directories]

    def find_documents(self, modified_since=None):
        documents = []
        for directory in self._directories:
            if not directory.exists():
                logger.warning(f"Word doc directory not found: {directory}")
                continue
            if not directory.is_dir():
                logger.warning(f"Path is not a directory: {directory}")
                continue

            for pattern in _SUPPORTED_EXTENSIONS:
                for file_path in directory.rglob(pattern):
                    if file_path.name.startswith('~$'):
                        continue
                    try:
                        stat_result = file_path.stat()
                        mod_time = datetime.fromtimestamp(stat_result.st_mtime)
                    except OSError as e:
                        logger.warning(f"Cannot stat {file_path}: {e}")
                        continue

                    if modified_since:
                        compare_date = modified_since
                        if isinstance(compare_date, datetime):
                            compare_date = compare_date.date()
                        if mod_time.date() < compare_date:
                            continue

                    documents.append({
                        'file_path': str(file_path),
                        'file_name': file_path.name,
                        'modified_at': mod_time.isoformat(),
                        'size_bytes': stat_result.st_size,
                    })

        logger.info(f"Found {len(documents)} document(s) (.docx + .pdf).")
        return documents

    def extract_text(self, file_path: str) -> dict:
        ext = Path(file_path).suffix.lower()
        if ext == '.pdf':
            return self._extract_pdf_text(file_path)
        return self._extract_docx_text(file_path)

    def _extract_docx_text(self, file_path: str) -> dict:
        try:
            doc = Document(file_path)
        except PackageNotFoundError:
            logger.error(f"Cannot open docx file: {file_path}")
            return {'error': f'Cannot open: {file_path}', 'content': []}
        except Exception as e:
            logger.error(f"Error reading {file_path}: {e}")
            return {'error': str(e), 'content': []}

        content_blocks = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                content_blocks.append({
                    'type': 'paragraph',
                    'text': text,
                })

        for table in doc.tables:
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(cells)
            if rows:
                content_blocks.append({
                    'type': 'table',
                    'rows': rows,
                })

        return {
            'file_name': Path(file_path).name,
            'file_path': file_path,
            'content': content_blocks,
        }

    def _extract_pdf_text(self, file_path: str) -> dict:
        try:
            from PyPDF2 import PdfReader
        except ImportError:
            logger.error("PyPDF2 not installed — cannot read PDF files.")
            return {'error': 'PyPDF2 not installed', 'content': []}

        try:
            reader = PdfReader(file_path)
        except Exception as e:
            logger.error(f"Cannot open PDF file {file_path}: {e}")
            return {'error': str(e), 'content': []}

        content_blocks = []
        for i, page in enumerate(reader.pages):
            try:
                text = page.extract_text()
                if text and text.strip():
                    content_blocks.append({
                        'type': 'paragraph',
                        'text': text.strip(),
                    })
            except Exception as e:
                logger.debug(f"Could not extract text from page {i+1} of {file_path}: {e}")

        return {
            'file_name': Path(file_path).name,
            'file_path': file_path,
            'content': content_blocks,
        }
