import os
import tempfile
from datetime import date, datetime
from pathlib import Path
from unittest.mock import MagicMock, patch

from django.test import TestCase
from docx import Document

from core.services.word_service import WordService


class WordServiceTest(TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()

    def test_find_documents_empty_dir(self):
        service = WordService([self.tmpdir])
        docs = service.find_documents()
        self.assertEqual(docs, [])

    def test_find_documents_with_docx(self):
        doc = Document()
        doc.add_paragraph('Test content')
        path = os.path.join(self.tmpdir, 'test.docx')
        doc.save(path)

        service = WordService([self.tmpdir])
        docs = service.find_documents()
        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0]['file_name'], 'test.docx')

    def test_skip_temp_files(self):
        doc = Document()
        doc.add_paragraph('Test')
        doc.save(os.path.join(self.tmpdir, '~$temp.docx'))
        doc.save(os.path.join(self.tmpdir, 'real.docx'))

        service = WordService([self.tmpdir])
        docs = service.find_documents()
        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0]['file_name'], 'real.docx')

    def test_extract_text(self):
        doc = Document()
        doc.add_paragraph('Hello World')
        doc.add_paragraph('Second paragraph')
        path = os.path.join(self.tmpdir, 'test.docx')
        doc.save(path)

        service = WordService([self.tmpdir])
        result = service.extract_text(path)
        self.assertEqual(len(result['content']), 2)
        self.assertEqual(result['content'][0]['text'], 'Hello World')
        self.assertEqual(result['content'][1]['text'], 'Second paragraph')

    def test_extract_text_invalid_file(self):
        path = os.path.join(self.tmpdir, 'notreal.docx')
        service = WordService([self.tmpdir])
        result = service.extract_text(path)
        self.assertIn('error', result)

    def test_nonexistent_directory(self):
        service = WordService(['/nonexistent/path/abc123'])
        docs = service.find_documents()
        self.assertEqual(docs, [])

    def test_find_documents_with_date_filter(self):
        doc = Document()
        doc.add_paragraph('Old content')
        path = os.path.join(self.tmpdir, 'old.docx')
        doc.save(path)
        # Set modification time to the past
        old_time = datetime(2020, 1, 1).timestamp()
        os.utime(path, (old_time, old_time))

        doc2 = Document()
        doc2.add_paragraph('New content')
        path2 = os.path.join(self.tmpdir, 'new.docx')
        doc2.save(path2)

        service = WordService([self.tmpdir])
        docs = service.find_documents(modified_since=date.today())
        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0]['file_name'], 'new.docx')


class GraphCalendarServiceTest(TestCase):
    def test_extract_teams_url(self):
        from core.services.graph_calendar_service import GraphCalendarService

        body_with_url = (
            "Join the meeting:\n"
            "https://teams.microsoft.com/l/meetup-join/19%3ameeting_abc123/0\n"
            "Or call in by phone..."
        )
        url = GraphCalendarService._extract_teams_url(body_with_url)
        self.assertIn('teams.microsoft.com', url)

    def test_extract_teams_url_no_match(self):
        from core.services.graph_calendar_service import GraphCalendarService

        self.assertEqual(GraphCalendarService._extract_teams_url('No teams link here'), '')
        self.assertEqual(GraphCalendarService._extract_teams_url(''), '')


class ConfigServiceTest(TestCase):
    def test_default_config(self):
        from core.services.config_service import ConfigService
        ConfigService.reset()
        config = ConfigService.get_instance()
        self.assertTrue(config.outlook_enabled)
        self.assertTrue(config.onenote_enabled)
        self.assertIsInstance(config.word_doc_directories, list)
